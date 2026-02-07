import argparse
import queue
import sys
import threading
import time
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from selenium import webdriver
from selenium.common.exceptions import TimeoutException, WebDriverException
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait


BASE_URL = "https://tsdh.hcmus.edu.vn/ketqua-xettuyen"
OUTPUT_COLUMNS = [
    "MSSV",
    "Họ và tên",
    "Giới tính",
    "Ngày sinh",
    "SBD",
    "CCCD",
    "Quê quán",
    "Nv",
    "Ngành/Nhóm ngành",
    "PTXT",
    "Kết quả xét tuyển",
    "Ghi chú",
]


class ExcelWriter:
    def __init__(self, path: Path) -> None:
        self.path = path
        if path.exists():
            self.wb = load_workbook(path)
            self.ws = self.wb.active
            self.written = max(0, self.ws.max_row - 1)
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "KET_QUA"
            self.ws.append(OUTPUT_COLUMNS)
            self.wb.save(path)
            self.written = 0

    def append_rows(self, rows: List[Dict[str, str]]) -> None:
        if not rows:
            return
        for row in rows:
            self.ws.append([row.get(column, "") for column in OUTPUT_COLUMNS])
        self.wb.save(self.path)
        self.written += len(rows)


def chunk_list(values: List[str], chunk_size: int) -> Iterable[List[str]]:
    for start in range(0, len(values), chunk_size):
        yield values[start : start + chunk_size]


def build_driver(chromedriver: Path, *, headless: bool = True) -> webdriver.Chrome:
    options = Options()
    if headless:
        options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-gpu")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--window-size=1280,900")

    service = Service(executable_path=str(chromedriver.resolve()))
    return webdriver.Chrome(service=service, options=options)


def parse_rows(html: str) -> List[Dict[str, str]]:
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find("table", class_="table")
    if not table:
        return []

    body = table.find("tbody")
    if not body:
        return []

    rows: List[Dict[str, str]] = []
    for tr in body.find_all("tr"):
        cells = tr.find_all("td")
        if not cells:
            continue

        if len(cells) == 1:
            message = cells[0].get_text(strip=True)
            if "Không tìm thấy kết quả tra cứu" in message:
                return []
            continue

        mapped: Dict[str, str] = {column: "" for column in OUTPUT_COLUMNS}
        mapped["SBD"] = cells[0].get_text(strip=True)
        mapped["CCCD"] = cells[1].get_text(strip=True)
        mapped["MSSV"] = cells[2].get_text(strip=True)
        if not mapped["MSSV"]:
            continue
        mapped["Họ và tên"] = cells[3].get_text(strip=True)
        mapped["Nv"] = cells[4].get_text(strip=True)
        mapped["Ngành/Nhóm ngành"] = cells[5].get_text(strip=True)
        mapped["Kết quả xét tuyển"] = cells[6].get_text(strip=True) or "XXXX"
        mapped["PTXT"] = cells[7].get_text(strip=True)
        mapped["Ghi chú"] = cells[8].get_text(strip=True)

        # Các cột không có trên giao diện được để trống.
        rows.append(mapped)

    return rows


def fetch_sbd_rows(driver: webdriver.Chrome, sbd: str, *, wait_timeout: int = 15) -> List[Dict[str, str]]:
    wait = WebDriverWait(driver, wait_timeout)

    driver.get(BASE_URL)
    query_input = wait.until(EC.presence_of_element_located((By.ID, "Query")))
    query_input.clear()
    query_input.send_keys(sbd)

    submit_btn = wait.until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, "input[type='submit']"))
    )
    submit_btn.click()

    try:
        wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "table.table tbody tr")))
    except TimeoutException:
        return []

    return parse_rows(driver.page_source)


def read_sbd_values(path: Path, limit: Optional[int]) -> Iterable[str]:
    with path.open() as handle:
        for idx, line in enumerate(handle):
            value = line.strip()
            if not value:
                continue
            yield value
            if limit is not None and idx + 1 >= limit:
                break


def parse_args(argv: List[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Tra cứu kết quả xét tuyển bằng headless Selenium và ghi ra Excel.")
    parser.add_argument("--sbd-file", type=Path, default=Path("sbd.txt"), help="Đường dẫn tới file chứa danh sách SBD.")
    parser.add_argument("--output", type=Path, default=Path("ketqua.xlsx"), help="Đường dẫn file Excel đầu ra.")
    parser.add_argument("--chromedriver", type=Path, default=Path("./chromedriver"), help="Đường dẫn tới chromedriver.")
    parser.add_argument("--limit", type=int, default=None, help="Giới hạn số lượng SBD (để thử nghiệm).")
    parser.add_argument("--throttle", type=float, default=0.3, help="Độ trễ giữa các lần tra cứu (giây).")
    parser.add_argument(
        "--workers",
        type=int,
        default=1,
        help="Số lượng worker chạy song song (mỗi worker giữ một driver riêng).",
    )
    parser.add_argument(
        "--chunk-size",
        type=int,
        default=100,
        help="Số lượng SBD tối đa mà mỗi worker lấy trong một lần (tối thiểu 1).",
    )
    return parser.parse_args(argv)


def main(argv: List[str]) -> int:
    args = parse_args(argv)

    if not args.sbd_file.exists():
        print(f"[!] Không tìm thấy file {args.sbd_file}", file=sys.stderr)
        return 1

    if not args.chromedriver.exists():
        print(f"[!] Không tìm thấy chromedriver tại {args.chromedriver}", file=sys.stderr)
        return 1
    if args.chunk_size <= 0:
        print("[!] chunk-size phải >= 1.", file=sys.stderr)
        return 1

    sbd_values = list(read_sbd_values(args.sbd_file, args.limit))
    if not sbd_values:
        print("[!] Không có SBD nào để tra cứu.", file=sys.stderr)
        return 1

    writer = ExcelWriter(args.output)

    if args.workers <= 1:
        collected = run_sequential(args, sbd_values, writer)
    else:
        collected = run_parallel(args, sbd_values, writer)

    if not collected:
        print("[!] Không thu được dữ liệu nào, dừng lại.", file=sys.stderr)
        return 1

    print(f"[✓] Đã ghi {writer.written} dòng vào {args.output}.")
    return 0


def run_sequential(
    args: argparse.Namespace, sbd_values: List[str], writer: ExcelWriter
) -> List[Dict[str, str]]:
    driver = build_driver(args.chromedriver, headless=True)
    collected: List[Dict[str, str]] = []
    try:
        for index, sbd in enumerate(sbd_values, start=1):
            rows = fetch_sbd_rows(driver, sbd)
            collected.extend(rows)
            writer.append_rows(rows)

            if rows:
                print(f"[+] {sbd}: {len(rows)} dòng.")

            if index % 5000 == 0:
                print(f"[*] Đã xử lý {index} / {len(sbd_values)} SBD.")

            if args.throttle > 0:
                time.sleep(args.throttle)
    finally:
        driver.quit()
    return collected


def run_parallel(
    args: argparse.Namespace, sbd_values: List[str], writer: ExcelWriter
) -> List[Dict[str, str]]:
    total = len(sbd_values)
    task_queue: "queue.Queue[Optional[List[str]]]" = queue.Queue()
    result_queue: "queue.Queue[Tuple[str, List[Dict[str, str]], Optional[Exception]]]" = queue.Queue()

    workers: List[threading.Thread] = []
    for idx in range(max(1, args.workers)):
        thread = threading.Thread(
            target=worker_thread,
            args=(idx, args.chromedriver, args.throttle, task_queue, result_queue),
            daemon=True,
        )
        thread.start()
        workers.append(thread)

    for chunk in chunk_list(sbd_values, args.chunk_size):
        task_queue.put(chunk)
    for _ in workers:
        task_queue.put(None)

    collected: List[Dict[str, str]] = []
    processed = 0
    errors = 0

    while processed < total:
        sbd, rows, error = result_queue.get()
        processed += 1

        collected.extend(rows)
        writer.append_rows(rows)
        if rows:
            print(f"[+] {sbd}: {len(rows)} dòng.")

        if error:
            errors += 1
            print(f"[!] Worker báo lỗi với {sbd}: {error}", file=sys.stderr)

        if processed % 6000 == 0:
            print(f"[*] Đã xử lý {processed} / {total} SBD.")

        result_queue.task_done()

    task_queue.join()
    result_queue.join()
    for thread in workers:
        thread.join(timeout=2)

    if errors:
        print(f"[!] Hoàn tất với {errors} lỗi (đã bỏ qua).", file=sys.stderr)

    return collected


def worker_thread(
    worker_id: int,
    chromedriver: Path,
    throttle: float,
    task_queue: "queue.Queue[Optional[List[str]]]",
    result_queue: "queue.Queue[Tuple[str, List[Dict[str, str]], Optional[Exception]]]",
) -> None:
    try:
        driver = build_driver(chromedriver, headless=True)
    except Exception as exc:
        result_queue.put(("", [], exc))
        return

    try:
        while True:
            chunk = task_queue.get()
            if chunk is None:
                task_queue.task_done()
                break

            for sbd in chunk:
                rows: List[Dict[str, str]] = []
                error: Optional[Exception] = None

                try:
                    rows = fetch_sbd_rows(driver, sbd)
                except WebDriverException as exc:
                    error = exc
                except Exception as exc:  # pragma: no cover - defensive
                    error = exc

                result_queue.put((sbd, rows, error))

                if throttle > 0:
                    time.sleep(throttle)

            task_queue.task_done()
    finally:
        driver.quit()


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))